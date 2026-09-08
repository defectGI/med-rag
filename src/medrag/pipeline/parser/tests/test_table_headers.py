"""`TableData.header_rows` (IR v6): which leading rows are header rows.

Deterministic sources fill it at parse time (html <th>/<thead>, markdown's pipe
header, docx w:tblHeader); tables/header_infer.py fills the rest via a
validated LLM answer marked with a "header-llm" flag; render/markdown.py stops
promoting a data row to the GFM header when a table is known headerless.
"""

from __future__ import annotations

import docx
import pytest
from docx.oxml import parse_xml

from medrag.pipeline.parser.parsers.base import (
    ParsedDocument,
    TableBlock,
    TableData,
    text_cell,
)
from medrag.pipeline.parser.parsers.docx_parser import DocxParser
from medrag.pipeline.parser.parsers.html_parser import HtmlParser
from medrag.pipeline.parser.parsers.markdown_parser import MarkdownParser
from medrag.pipeline.parser.render.markdown import _render_gfm, _render_html_table
from medrag.pipeline.parser.tables.header_infer import infer_headers

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


@pytest.fixture(autouse=True)
def _header_llm_default_on(monkeypatch):
    """Isolate from a developer's real TABLE_HEADER_LLM setting."""
    monkeypatch.delenv("TABLE_HEADER_LLM", raising=False)


# --- helpers -----------------------------------------------------------------


def _parse_html(tmp_path, body: str) -> ParsedDocument:
    p = tmp_path / "t.html"
    p.write_text(f"<html><body>{body}</body></html>", encoding="utf-8")
    return HtmlParser().parse(p, "t")


def _table(rows: list[list[str]], header_rows: int | None = None) -> TableData:
    return TableData(n_rows=len(rows), n_cols=len(rows[0]),
                     cells=[[text_cell(v) for v in r] for r in rows],
                     header_rows=header_rows)


def _doc_with(table: TableData) -> ParsedDocument:
    doc = ParsedDocument(doc_id="d", source_path="x", fmt="test")
    doc.blocks = [TableBlock(id="t0", table=table)]
    return doc


class FakeLLM:
    def __init__(self, reply: str) -> None:
        self.reply = reply
        self.calls = 0

    def complete(self, *, system, user, max_tokens=1024):
        self.calls += 1
        return self.reply


# --- serialization -----------------------------------------------------------


def test_header_rows_roundtrips_and_is_absent_when_unknown():
    known = _table([["H", "H2"], ["a", "b"]], header_rows=1)
    assert known.to_dict()["header_rows"] == 1
    assert TableData.from_dict(known.to_dict()).header_rows == 1

    unknown = _table([["a", "b"]])
    assert "header_rows" not in unknown.to_dict()
    assert TableData.from_dict(unknown.to_dict()).header_rows is None

    headerless = _table([["a", "b"]], header_rows=0)
    assert headerless.to_dict()["header_rows"] == 0  # 0 is a fact, not "unset"
    assert TableData.from_dict(headerless.to_dict()).header_rows == 0


# --- html --------------------------------------------------------------------


def test_html_thead_sets_header_rows(tmp_path):
    doc = _parse_html(tmp_path, (
        "<table><thead><tr><td>Name</td><td>Min</td></tr></thead>"
        "<tbody><tr><td>a</td><td>1</td></tr></tbody></table>"))
    assert doc.tables()[0].table.header_rows == 1


def test_html_all_th_rows_set_header_rows(tmp_path):
    doc = _parse_html(tmp_path, (
        "<table>"
        "<tr><th>Group</th><th>Group</th></tr>"
        "<tr><th>Min</th><th>Max</th></tr>"
        "<tr><td>1</td><td>2</td></tr>"
        "</table>"))
    assert doc.tables()[0].table.header_rows == 2


def test_html_all_td_table_stays_unknown(tmp_path):
    doc = _parse_html(tmp_path, (
        "<table><tr><td>Name</td><td>Min</td></tr>"
        "<tr><td>a</td><td>1</td></tr></table>"))
    assert doc.tables()[0].table.header_rows is None


def test_html_row_header_column_is_not_a_header_row(tmp_path):
    # <th> only in column 0 (row labels) must not claim leading header ROWS
    doc = _parse_html(tmp_path, (
        "<table><tr><th>Speed</th><td>10</td></tr>"
        "<tr><th>Weight</th><td>3</td></tr></table>"))
    assert doc.tables()[0].table.header_rows is None


def test_html_nested_tables_track_headers_independently(tmp_path):
    doc = _parse_html(tmp_path, (
        "<table><tr><th>Outer</th></tr><tr><td>"
        "<table><tr><td>x</td></tr></table>"
        "</td></tr></table>"))
    outer = doc.tables()[0].table
    assert outer.header_rows == 1
    inner = next(b for b in outer.cells[1][0].blocks if isinstance(b, TableBlock))
    assert inner.table.header_rows is None


# --- markdown / docx ---------------------------------------------------------


def test_markdown_pipe_table_header_is_explicit(tmp_path):
    p = tmp_path / "t.md"
    p.write_text("| Name | Min |\n| --- | --- |\n| a | 1 |\n", encoding="utf-8")
    doc = MarkdownParser().parse(p, "t")
    assert doc.tables()[0].table.header_rows == 1


def test_docx_tbl_header_property(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Name"
    t.cell(1, 0).text = "a"
    tr0 = t.rows[0]._tr
    tr0.insert(0, parse_xml(f"<w:trPr {_W}><w:tblHeader/></w:trPr>"))
    p = tmp_path / "t.docx"
    d.save(str(p))
    doc = DocxParser().parse(p, "t")
    assert doc.tables()[0].table.header_rows == 1


def test_docx_without_tbl_header_stays_unknown(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Name"
    p = tmp_path / "t.docx"
    d.save(str(p))
    doc = DocxParser().parse(p, "t")
    assert doc.tables()[0].table.header_rows is None


# --- rendering ---------------------------------------------------------------


def test_gfm_headerless_table_gets_an_empty_header_row():
    md = _render_gfm(_table([["a", "1"], ["b", "2"]], header_rows=0))
    lines = md.split("\n")
    assert lines[0] == "|  |  |"          # no data row promoted to header
    assert lines[1].count("---") == 2
    assert "| a | 1 |" in lines and "| b | 2 |" in lines


def test_gfm_unknown_header_keeps_first_row_promotion():
    md = _render_gfm(_table([["Name", "Min"], ["a", "1"]]))
    assert md.split("\n")[0] == "| Name | Min |"


def test_html_render_marks_header_rows_as_th():
    html = _render_html_table(_table([["Name", "Min"], ["a", "1"]], header_rows=1))
    assert "<th>Name</th><th>Min</th>" in html
    assert "<td>a</td><td>1</td>" in html


def test_html_render_without_header_info_uses_td_everywhere():
    html = _render_html_table(_table([["Name", "Min"], ["a", "1"]]))
    assert "<th>" not in html


# --- LLM fallback ------------------------------------------------------------


def test_infer_sets_header_rows_and_marks_provenance():
    doc = _doc_with(_table([["Name", "Min"], ["a", "1"], ["b", "2"]]))
    llm = FakeLLM("1")
    infer_headers(doc, client=llm)
    block = doc.tables()[0]
    assert block.table.header_rows == 1
    assert "header-llm" in (block.table_flags or [])
    assert llm.calls == 1


def test_infer_accepts_zero_as_headerless():
    doc = _doc_with(_table([["a", "1"], ["b", "2"]]))
    infer_headers(doc, client=FakeLLM("0"))
    block = doc.tables()[0]
    assert block.table.header_rows == 0
    assert "header-llm" in (block.table_flags or [])


def test_infer_never_touches_a_deterministic_value():
    doc = _doc_with(_table([["Name", "Min"], ["a", "1"]], header_rows=1))
    llm = FakeLLM("0")
    infer_headers(doc, client=llm)
    assert llm.calls == 0
    assert doc.tables()[0].table.header_rows == 1
    assert not doc.tables()[0].table_flags


def test_infer_rejects_out_of_range_answers():
    # 2 rows: an answer of 2 would leave no data row -> misread, keep None
    doc = _doc_with(_table([["Name", "Min"], ["a", "1"]]))
    infer_headers(doc, client=FakeLLM("2"))
    assert doc.tables()[0].table.header_rows is None
    assert not doc.tables()[0].table_flags


def test_infer_survives_a_dead_client():
    class DeadLLM:
        def complete(self, **kw):
            raise RuntimeError("connection refused")

    doc = _doc_with(_table([["Name", "Min"], ["a", "1"]]))
    infer_headers(doc, client=DeadLLM())
    assert doc.tables()[0].table.header_rows is None


def test_infer_kill_switch(monkeypatch):
    monkeypatch.setenv("TABLE_HEADER_LLM", "0")
    doc = _doc_with(_table([["Name", "Min"], ["a", "1"]]))
    llm = FakeLLM("1")
    infer_headers(doc, client=llm)
    assert llm.calls == 0
    assert doc.tables()[0].table.header_rows is None
